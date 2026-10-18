import asyncio
import base64
import threading
import time

import pytest
from fastapi.testclient import TestClient

from app.actions import Actions
from app.agents import BrainResult, ungrounded_values
from app.config import Settings
from app.main import create_app
from app.meeting import EMPTY_WAKE_SECONDS, WAKE_PROMPT, MeetingRuntime
from app.store import make_store


class FakeBQ:
    def __init__(self):
        self.rows = []

    def insert_rows_json(self, table, rows):
        self.rows += rows
        return []


class FakeBrain:
    def __init__(self, actions):
        self.actions = actions
        self.calls = []

    async def handle(self, meeting_id, requester, command, context, company=None, **hints):
        self.calls.append(command)
        text = command.lower()
        if "jira" in text or "tickets" in text:
            card = await self.actions.draft_jira_plan(meeting_id, requester, "SEC", "Compliance fixes", [
                dict(title="Disable TLS 1.1 on legacy-settle", story_points=3, priority="Medium", assignee="David"),
                dict(title="Retain audit logs 365 days", story_points=5, priority="Medium", assignee="David")], ["chat_upload:spec.pdf"])
            return BrainResult("Drafted two tickets.", "", cards_created=[card["card_id"]])
        if "change the first" in text:
            card = (await self.actions.store.cards(meeting_id))[-1]
            await self.actions.update_jira_task(meeting_id, card["card_id"], requester, 1, priority="High")
            return BrainResult("Updated.", "", cards_updated=[card["card_id"]])
        if "conflict" in text:
            return BrainResult("Found 2 compliance conflicts in the PDF.", "| Rule | Spec |\n|---|---|\n| TLS 1.2+ | TLS 1.1 |")
        return BrainResult("Payment success rate was 98.9 percent.", "")


@pytest.fixture
def env():
    settings = Settings(store="memory", ai_enabled=False)
    store = make_store("memory", settings.project)
    bq = FakeBQ()
    actions = Actions(settings, store, bq_client=bq)
    brain = FakeBrain(actions)
    rt = MeetingRuntime("m-1", store, brain, actions)
    events = []

    async def listener(event, data):
        events.append((event, data))
    store.subscribe("m-1", listener)
    return rt, store, actions, brain, bq, events


async def test_wake_from_interim_then_command(env):
    rt, store, _, brain, _, events = env
    await rt.on_speech("Vyshnav", "hey tensor what", final=False)
    assert (await store.tensor("m-1"))["status"] == "listening"
    await rt.on_speech("Vyshnav", "Hey Tensor, what was our payment success rate yesterday?", final=True)
    assert brain.calls == ["what was our payment success rate yesterday?"]
    speak = [d for e, d in events if e == "speak"][-1]
    assert speak["text"] == "Payment success rate was 98.9 percent."


async def test_empty_wake_waits_for_the_same_speaker(env):
    rt, _, _, brain, _, _ = env
    await rt.on_utterance("Vyshnav", "Hey Tensor.", "voice")
    await rt.on_utterance("Sarah", "I think we should wait", "voice")
    assert brain.calls == []
    await rt.on_utterance("Vyshnav", "what was the success rate", "voice")
    assert brain.calls == ["what was the success rate"]


async def test_reported_repeated_names_acknowledge_once_then_accept_a_question(env):
    rt, store, _, brain, _, events = env

    class Voice:
        def synthesize(self, text):
            return b"test audio"

    rt.voice = Voice()
    for text in ["Hey, tensor.", "Tensor.", "Tensor."]:
        await rt.on_speech("Bot-1", text, final=True)
    tensor = await store.tensor("m-1")
    assert tensor["status"] == "listening" and tensor["awaiting_command"] is True
    assert brain.calls == []
    acknowledgements = [d for e, d in events if e == "speak"]
    assert len(acknowledgements) == 1
    assert acknowledgements[0]["text"] == WAKE_PROMPT
    assert base64.b64decode(acknowledgements[0]["audio_mp3_base64"]) == b"test audio"
    assert [m["text"] for m in await store.chat("m-1")] == [WAKE_PROMPT]
    await rt.on_speech("Bot-2", "We should discuss the launch", final=True)
    assert brain.calls == []
    await rt.on_speech("Bot-1", "what was our success rate?", final=True)
    assert brain.calls == ["what was our success rate?"]
    assert (await store.tensor("m-1"))["awaiting_command"] is False


async def test_greeting_and_name_in_separate_segments(env):
    rt, store, _, brain, _, _ = env
    for text in ["Hey.", "Tensor.", "What were our active users?"]:
        await rt.on_speech("Bot-1", text, final=True)
    assert brain.calls == ["What were our active users?"]
    assert len(await store.transcript("m-1")) == 3


async def test_empty_wake_timeout_does_not_consume_later_conversation(env, monkeypatch):
    rt, store, _, brain, _, _ = env
    clock = [time.time()]
    monkeypatch.setattr("app.meeting.time.time", lambda: clock[0])
    await rt.on_utterance("Bot-1", "Tensor.", "voice")
    waiting = rt.awaiting_command
    clock[0] += EMPTY_WAKE_SECONDS + 1
    await rt._expire_wake(waiting)
    assert (await store.tensor("m-1"))["status"] == "idle"
    assert rt.awaiting_command is None
    await rt.on_utterance("Bot-1", "Let's discuss something else", "voice")
    assert brain.calls == []


async def test_repeated_name_refreshes_the_window_without_duplicate_greetings(env, monkeypatch):
    rt, store, _, brain, _, events = env
    clock = [time.time()]
    monkeypatch.setattr("app.meeting.time.time", lambda: clock[0])
    await rt.on_utterance("Bot-1", "Hey, Tensor.", "voice")
    old_waiting = rt.awaiting_command
    clock[0] += EMPTY_WAKE_SECONDS - 5
    await rt.on_utterance("Bot-1", "Tensor.", "voice")
    await rt._expire_wake(old_waiting)  # an older timeout cannot close the refreshed window
    clock[0] += 10
    await rt.on_utterance("Bot-1", "What was the success rate?", "voice")
    assert brain.calls == ["What was the success rate?"]
    assert len([d for e, d in events if e == "speak" and d["text"] == WAKE_PROMPT]) == 1


async def test_wake_stays_quiet_while_someone_else_is_talking(env):
    rt, store, _, brain, _, events = env
    await rt.vad("Bot-2", True)
    await rt.on_utterance("Bot-1", "Hey, Tensor.", "voice")
    assert (await store.tensor("m-1"))["awaiting_command"] is True
    assert [m["text"] for m in await store.chat("m-1")] == [WAKE_PROMPT]
    assert not [d for e, d in events if e == "speak"]
    assert brain.calls == []


async def test_wake_audio_does_not_start_over_the_callers_next_words(env):
    rt, store, _, _, _, events = env
    await rt.vad("Bot-1", True)
    await rt.on_utterance("Bot-1", "Tensor.", "voice")
    assert (await store.tensor("m-1"))["awaiting_command"] is True
    assert not [d for e, d in events if e == "speak"]
    await rt.vad("Bot-1", False)
    assert "Bot-1" not in rt.active_speakers


def test_long_running_speech_remains_active_until_vad_stops(env):
    rt, _, _, _, _, _ = env
    rt.active_speakers.add("Bot-2")
    rt.voice_activity["Bot-2"] = time.time() - 10
    assert rt.humans_speaking()
    assert not rt.humans_speaking(exclude="Bot-2")
    rt.close_participant("Bot-2")
    assert not rt.humans_speaking()


async def test_barge_in_on_greeting_keeps_the_question_window_open(env):
    rt, store, _, brain, _, events = env
    await rt.on_utterance("Bot-1", "Tensor.", "voice")
    await rt.vad("Bot-1", True)
    assert (await store.tensor("m-1"))["status"] == "listening"
    assert any(e == "stop_audio" for e, _ in events)
    await rt.on_utterance("Bot-1", "What was the success rate?", "voice")
    assert brain.calls == ["What was the success rate?"]


async def test_slow_wake_audio_cannot_interrupt_the_answer(env):
    rt, _, _, brain, _, events = env
    started, release = threading.Event(), threading.Event()

    class SlowWakeVoice:
        def synthesize(self, text):
            if text == WAKE_PROMPT:
                started.set()
                release.wait(timeout=5)
            return b"test audio"

    rt.voice = SlowWakeVoice()
    greeting = asyncio.create_task(rt.on_utterance("Bot-1", "Tensor.", "voice"))
    try:
        assert await asyncio.to_thread(started.wait, 5)
        await rt.on_utterance("Bot-1", "What was the success rate?", "voice")
    finally:
        release.set()
        await greeting
    assert brain.calls == ["What was the success rate?"]
    assert not [d for e, d in events if e == "speak" and d["text"] == WAKE_PROMPT]


async def test_repeated_name_does_not_replace_thinking_state(env):
    rt, store, _, brain, _, events = env
    await store.set_tensor("m-1", status="thinking")
    async with rt.lock:
        await rt.on_utterance("Bot-1", "Tensor.", "voice")
    assert (await store.tensor("m-1"))["status"] == "thinking"
    assert rt.awaiting_command is None
    assert not [d for e, d in events if e == "speak"]


async def test_card_prompt_then_pure_voice_approval_executes_once_without_the_model(env):
    rt, store, _, brain, bq, events = env
    await rt.on_utterance("Vyshnav", "Hey Tensor, turn those into Jira tickets for David", "voice")
    card = (await store.cards("m-1"))[0]
    assert card["status"] == "PENDING_REVIEW"
    assert [d for e, d in events if e == "speak"][-1]["text"].startswith("I've put the draft on screen")
    await rt.on_utterance("Vyshnav", "Looks good, go ahead", "voice")  # no wake word: follow-up window
    card = (await store.cards("m-1"))[0]
    assert card["status"] == "EXECUTED" and card["result"]["refs"] == ["SEC-401", "SEC-402"]
    assert card["approved_by"] == "Vyshnav" and card["approval_channel"] == "voice:Vyshnav"
    assert len(brain.calls) == 1  # approval decided in code
    assert bq.rows[0]["external_refs"] == ["SEC-401", "SEC-402"] and bq.rows[0]["approved_by"] == "Vyshnav"
    with pytest.raises(ValueError):
        await rt.actions.approve("m-1", card["card_id"], "Sarah", "ui")
    with pytest.raises(ValueError):
        await rt.actions.execute("m-1", card["card_id"])


async def test_edit_then_approve_in_one_sentence(env):
    rt, store, _, brain, _, _ = env
    await rt.on_utterance("Vyshnav", "Hey Tensor, plan Jira tickets for those conflicts", "voice")
    await rt.on_utterance("Vyshnav", "Change the first one to High priority and go ahead", "voice")
    card = (await store.cards("m-1"))[0]
    assert card["payload"]["proposed_tasks"][0]["priority"] == "High"
    assert card["status"] == "EXECUTED" and [h["change"] for h in card["history"]][-1] == "task 1: priority=High"


async def test_followup_needs_a_pending_card_and_card_language(env):
    rt, _, _, brain, _, _ = env
    rt.followup_until = time.time() + 60
    await rt.on_utterance("Sarah", "go ahead and share your screen", "voice")  # no pending card
    await rt.on_utterance("Sarah", "the weather is nice", "voice")
    assert brain.calls == []


async def test_long_answer_while_humans_talk_raises_hand_then_go_ahead_speaks(env):
    rt, store, _, _, _, events = env
    await rt.vad("Sarah", True)
    await rt.on_utterance("Vyshnav", "Hey Tensor, does this spec conflict with our security guidelines?", "voice")
    tensor = await store.tensor("m-1")
    assert tensor["status"] == "hand_raised" and tensor["hand"]["topic"] == "Found 2 compliance conflicts in the PDF."
    assert not [e for e, _ in events if e == "speak"]
    chat = [d for e, d in events if e == "chat"][-1]
    assert chat["author"] == "Tensor" and "TLS 1.1" in chat["text"] and chat["delivery"] == "HAND_RAISE"
    await rt.on_utterance("Priya", "Okay, go ahead Tensor", "voice")
    assert [d for e, d in events if e == "speak"][-1]["text"].startswith("Found 2 compliance conflicts")
    assert (await store.tensor("m-1")).get("hand") is None


async def test_barge_in_stops_tensor(env):
    rt, store, _, _, _, events = env
    await rt.speak("A short answer that takes a few seconds to say aloud.")
    await rt.vad("Sarah", True)
    assert ("stop_audio", dict(reason="barge_in", by="Sarah")) in events
    assert (await store.tensor("m-1"))["status"] == "idle"


async def test_chat_mention_and_cancel(env):
    rt, store, _, brain, _, _ = env
    await rt.on_chat("Priya", "@Tensor make jira tickets for this")
    assert brain.calls == ["make jira tickets for this"]
    await rt.on_utterance("Priya", "cancel that", "voice")
    assert (await store.cards("m-1"))[0]["status"] == "CANCELLED"


def test_grounding_check():
    assert ungrounded_values("Average 219,748 vs 217,441; p95 612 ms", [dict(a=219748.43, b=217441.28, c=612)]) == []
    assert ungrounded_values("about 218,600 users", [dict(a=219748.43)]) == ["218,600users"]


def seed_company(store, cid="acme", code="invite-acme-123"):
    team = [dict(key="david", name="David Mathew", role="Security Lead", email="david@acme.example", project="SEC")]
    asyncio.run(store.backend.put("companies", cid, dict(id=cid, name=cid.title(), industry="demo", jira_key="SEC", dataset=f"tensor_co_{cid}",
                                                        invite_code=code, status="READY", created_by=None, team=team, examples=[])))
    asyncio.run(store.backend.put("invites", code, dict(company_id=cid)))


def make_app():
    settings = Settings(store="memory", ai_enabled=False)
    store = make_store("memory", settings.project)
    actions = Actions(settings, store, bq_client=FakeBQ())
    return create_app(settings, store=store, brain=FakeBrain(actions), actions=actions, knowledge=object(), docs=object()), store


def test_websocket_room_snapshot_chat_and_rest():
    app, store = make_app()
    seed_company(store)
    with TestClient(app) as client:
        assert client.get("/api/health").json()["jira"] == "sandbox"
        assert client.get("/api/auth/me").status_code == 401
        assert client.get("/api/meetings/demo-room").status_code == 401
        with client.websocket_connect("/ws/demo-room") as ws:  # not signed in: accepted, then closed with 4401
            with pytest.raises(Exception):
                ws.receive_json()
        assert client.get("/api/invites/invite-acme-123").json()["company_name"] == "Acme"
        r = client.post("/api/auth/register", json=dict(name="Vyshnav", email="V@Example.com", password="correct horse", invite="invite-acme-123"))
        assert r.status_code == 200 and r.json()["user"]["role"] == "admin"  # first member of a seeded company runs it
        assert r.json()["company"]["invite_code"] == "invite-acme-123"
        with client.websocket_connect("/ws/demo-room") as ws:
            assert ws.receive_json()["event"] == "participants"
            snap = ws.receive_json()
            assert snap["event"] == "snapshot" and snap["data"]["meeting"]["id"] == "acme__demo-room"
            ws.send_json(dict(type="chat", text="@Tensor create jira tickets for the conflicts"))
            for _ in range(30):
                msg = ws.receive_json()
                if msg["event"] == "card":
                    card = msg["data"]
                    break
            assert card["status"] == "PENDING_REVIEW" and card["payload"]["proposed_tasks"][0]["assignee"] == "David Mathew"
            r = client.patch(f"/api/meetings/demo-room/cards/{card['card_id']}/tasks/2", json=dict(include=False))
            assert r.json()["payload"]["proposed_tasks"][1]["include"] is False
            r = client.post(f"/api/meetings/demo-room/cards/{card['card_id']}/approve", json={})
            assert r.json()["status"] == "EXECUTED" and r.json()["result"]["refs"] == ["SEC-401"]
            assert r.json()["approved_by"] == "Vyshnav"  # the session, not the request body, names the approver
            assert client.post(f"/api/meetings/demo-room/cards/{card['card_id']}/approve", json={}).status_code == 409
        assert client.post("/api/auth/logout", headers={"Origin": "https://evil.example"}).status_code == 403  # cross-site
        assert client.get("/kb/cloud-security-guidelines.html").status_code == 200
        assert client.get("/kb/../secret").status_code == 404
        assert client.get("/api/meetings/BAD_ID").status_code == 400


def test_companies_are_isolated_and_login_rules():
    app, store = make_app()
    seed_company(store, "acme", "invite-acme-123")
    seed_company(store, "globex", "invite-globex-456")
    with TestClient(app) as a, TestClient(app) as b:
        a.post("/api/auth/register", json=dict(name="Ann", email="ann@a.example", password="password-a", invite="invite-acme-123"))
        b.post("/api/auth/register", json=dict(name="Bob", email="bob@b.example", password="password-b", invite="invite-globex-456"))
        a.post("/api/meetings/standup/chat", json=dict(text="acme only"))
        import time as _t
        _t.sleep(0.1)
        assert [m["text"] for m in a.get("/api/meetings/standup").json()["chat"]] == ["acme only"]
        assert b.get("/api/meetings/standup").json()["chat"] == []  # same meeting name, different company
        r = b.post("/api/auth/register", json=dict(name="Eve", email="ANN@a.example", password="whatever1", invite="invite-globex-456"))
        assert r.status_code == 409  # emails are unique
        b.post("/api/auth/logout")
        assert b.get("/api/auth/me").status_code == 401
        assert b.post("/api/auth/login", json=dict(email="bob@b.example", password="wrong")).status_code == 401
        assert b.post("/api/auth/login", json=dict(email="BOB@b.example", password="password-b")).json()["company"]["id"] == "globex"
        assert "invite_code" in b.get("/api/auth/me").json()["company"]  # first member is admin
        joined = TestClient(app)
        joined.post("/api/auth/register", json=dict(name="Cat", email="cat@b.example", password="password-c", invite="invite-globex-456"))
        me = joined.get("/api/auth/me").json()
        assert me["user"]["role"] == "member" and "invite_code" not in me["company"]
        assert joined.post("/api/company/invite").status_code == 403  # admin only
        assert [x["name"] for x in joined.get("/api/directory").json()][-2:] == ["Bob", "Cat"]


def test_voice_signalling_reaches_only_the_named_participant_in_the_same_meeting():
    app, store = make_app()
    seed_company(store, "acme", "invite-acme-123")
    seed_company(store, "globex", "invite-globex-456")
    with TestClient(app) as client:  # one event loop, like the real server; each person has their own session cookie
        def session(name, email, invite):
            client.cookies.clear()
            client.post("/api/auth/register", json=dict(name=name, email=email, password="password-x", invite=invite))
            return {"cookie": f"tensor_session={client.cookies.get('tensor_session')}"}

        ann, bob = session("Ann", "ann@a.example", "invite-acme-123"), session("Bob", "bob@a.example", "invite-acme-123")
        other_bob = session("Bob", "bob@g.example", "invite-globex-456")
        client.cookies.clear()
        with client.websocket_connect("/ws/standup", headers=ann) as wa, client.websocket_connect("/ws/standup", headers=bob) as wb,                 client.websocket_connect("/ws/standup", headers=other_bob) as wc:
            for sock in (wa, wb, wc):  # like the app: signal only once the room snapshot has arrived
                while sock.receive_json()["event"] != "snapshot":
                    pass
            wa.send_json(dict(type="rtc", to="Bob", data=dict(description=dict(type="offer", sdp="v=0"))))
            wa.send_json(dict(type="mic", on=True))
            got = []
            for _ in range(10):
                msg = wb.receive_json()
                got.append(msg["event"])
                if msg["event"] == "rtc":
                    assert msg["data"] == {"from": "Ann", "data": {"description": {"type": "offer", "sdp": "v=0"}}}
                if msg["event"] == "presence":
                    assert msg["data"] == {"name": "Ann", "mic": True}
                    break
            assert "rtc" in got
            wc.send_json(dict(type="ping"))
            events = []
            for _ in range(5):
                events.append(wc.receive_json()["event"])
                if events[-1] == "pong":
                    break
            assert "rtc" not in events and "presence" not in events  # the other company's Bob hears nothing


async def test_streamed_voice_sends_audio_and_stops_on_barge_in(env):
    rt, _, _, _, _, events = env
    gate = threading.Event()

    class StreamingVoice:
        def stream(self, text, voice_id=None):
            yield b"\x01\x00" * 100  # first audio arrives quickly
            gate.wait(timeout=5)
            for _ in range(20):
                yield b"\x02\x00" * 6000

    rt.voice = StreamingVoice()
    await rt.speak("Payment success rate was 98.9 percent.")
    speak = [d for e, d in events if e == "speak"][-1]
    assert speak["stream"] is True and speak["audio_mp3_base64"] is None
    await asyncio.sleep(0.05)
    assert [d["seq"] for e, d in events if e == "audio"] == [0]  # the first chunk went out at once
    await rt.vad("Sarah", True)  # a human starts talking over Tensor
    gate.set()
    await asyncio.sleep(0.2)
    assert ("stop_audio" in [e for e, _ in events]) and len([1 for e, _ in events if e == "audio"]) == 1


async def test_streamed_voice_finishes_with_an_end_marker(env):
    rt, _, _, _, _, events = env

    class StreamingVoice:
        def stream(self, text, voice_id=None):
            for _ in range(3):
                yield b"\x01\x00" * 6000

    rt.voice = StreamingVoice()
    await rt.speak("Done.")
    for _ in range(50):
        await asyncio.sleep(0.01)
        if any(e == "audio" and d.get("end") for e, d in events):
            break
    audio = [d for e, d in events if e == "audio"]
    assert audio[-1]["end"] is True and sum(len(base64.b64decode(d["pcm"])) for d in audio) == 36000


def test_admin_ends_the_meeting_for_everyone():
    app, store = make_app()
    seed_company(store)
    with TestClient(app) as admin, TestClient(app) as member:
        admin.post("/api/auth/register", json=dict(name="Ann", email="ann@a.example", password="password-a", invite="invite-acme-123"))
        member.post("/api/auth/register", json=dict(name="Bob", email="bob@a.example", password="password-b", invite="invite-acme-123"))
        with member.websocket_connect("/ws/standup") as ws:
            ws.send_json(dict(type="chat", text="@Tensor create jira tickets for the conflicts"))
            for _ in range(30):
                if ws.receive_json()["event"] == "card":
                    break
            assert member.post("/api/meetings/standup/end", json={}).status_code == 403  # members can't
            assert admin.post("/api/meetings/standup/end", json={}).json()["ended_by"] == "Ann"
            seen = []
            with pytest.raises(Exception):
                for _ in range(30):
                    seen.append(ws.receive_json()["event"])
            assert "ended" in seen
        snap = admin.get("/api/meetings/standup").json()
        assert snap["meeting"]["ended_by"] == "Ann" and snap["transcript"] == []
        assert [c["status"] for c in snap["cards"]] == ["CANCELLED"]  # open drafts are cancelled, never executed
        with admin.websocket_connect("/ws/standup") as ws:  # joining again starts the meeting afresh
            ws.receive_json()
            assert ws.receive_json()["data"]["meeting"]["ended_at"] is None



async def test_live_email_only_reaches_registered_company_members(monkeypatch):
    import json as _json

    import httpx as _httpx

    settings = Settings(store="memory", ai_enabled=False, gmail_client_id="id", gmail_client_secret="s", gmail_refresh_token="r",
                        gmail_sender="tensor.sender@gmail.com")
    store = make_store("memory", settings.project)
    await store.backend.put("users", "u1", dict(uid="u1", email="ann@a.example", name="Ann", company_id="acme"))
    actions = Actions(settings, store, bq_client=FakeBQ())
    monkeypatch.setattr(actions, "_gmail_token", lambda: "token")
    sent = []

    class FakeHttp:
        def __init__(self, *a, **k):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *exc):
            return False

        async def post(self, url, headers=None, json=None):
            import base64 as b64
            sent.append(b64.urlsafe_b64decode(json["raw"]).decode())
            return _httpx.Response(200, json={"id": "gm-1"}, request=_httpx.Request("POST", url))

    monkeypatch.setattr("app.actions.httpx.AsyncClient", FakeHttp)
    people = [dict(key="ann", name="Ann", email="ann@a.example"), dict(key="david", name="David Mathew", email="david@acme.example")]
    card = await actions.draft_email("acme__m", "Ann", ["Ann"], ["David"], "Tickets", "Here they are.", [], people)
    done = await actions.approve("acme__m", card["card_id"], "Ann", "ui_click:Ann")
    assert done["status"] == "EXECUTED" and done["result"]["sent_to"] == ["ann@a.example"] and done["result"]["skipped"] == ["David Mathew"]
    assert "To: ann@a.example" in sent[0] and "david@acme.example" not in sent[0] and "Reply-To: Ann <ann@a.example>" in sent[0]
    stranger = await actions.draft_email("acme__m", "Ann", ["David"], [], "Hi", "x", [], people)
    failed = await actions.approve("acme__m", stranger["card_id"], "Ann", "ui_click:Ann")
    assert failed["status"] == "FAILED" and "registered member" in failed["error"] and len(sent) == 1


async def test_answer_is_spoken_while_the_details_are_still_being_written(env):
    rt, store, _, _, _, events = env
    spoke_before_done = []

    class StreamingBrain:
        async def handle(self, meeting_id, requester, command, context, company=None, on_partial=None, **hints):
            text = ""
            for piece in ["SPOKEN: Payment success ", "was 98.9 percent.\nDETAILS:\n", "| Day | Rate |\n|---|---|\n",
                          "| Oct 7 | 96.12% |\n", "| Oct 8 | 98.9% |\n"]:
                text += piece
                await on_partial(text, dict(cards=False, outputs=[dict(rate=0.989), dict(rate=0.9612)], command=command, context=context))
                await asyncio.sleep(0.25)
            spoke_before_done.append(any(e == "speak" for e, _ in events))
            return BrainResult("Payment success was 98.9 percent.", "| Day | Rate |\n|---|---|\n| Oct 7 | 96.12% |\n| Oct 8 | 98.9% |")

    rt.brain = StreamingBrain()
    await rt.on_chat("Vyshnav", "@Tensor what was payment success on Oct 7 and 8?")
    for _ in range(50):
        await asyncio.sleep(0.02)
    assert spoke_before_done == [True]  # Tensor talked before the reply was finished
    answers = [m for m in await store.chat("m-1") if m["author"] == "Tensor"]
    assert len(answers) == 1 and answers[0]["streaming"] is False and "Oct 8" in answers[0]["text"]
    assert len([d for e, d in events if e == "chat" and d["id"] == answers[0]["id"]]) >= 2  # updated in place
    assert len([d for e, d in events if e == "speak"]) == 1  # spoken once, not again at the end


def test_documents_are_downloadable_as_word_and_markdown():
    from io import BytesIO

    from docx import Document as Docx

    app, store = make_app()
    seed_company(store)
    actions = app.state.actions
    with TestClient(app) as client:
        client.post("/api/auth/register", json=dict(name="Ann", email="ann@a.example", password="password-a", invite="invite-acme-123"))
        md = ("# Meeting minutes: standup\n\n**Date:** Sunday 18 October 2026\n\n## Summary\nWe agreed to ship.\n\n"
              "## Action items\n| Action | Owner | Due | Ticket |\n|---|---|---|---|\n| Run load test | Rahul | Thu | SEC-401 |\n\n- one\n- **two**\n")
        card = asyncio.run(actions.create_document("acme__standup", "Ann", "Minutes: standup", "minutes", md))
        assert card["status"] == "READY" and card["type"] == "DOCUMENT"
        r = client.get(f"/api/meetings/standup/exports/{card['card_id']}.docx")
        assert r.status_code == 200 and "Minutes_standup.docx" in r.headers["content-disposition"]
        doc = Docx(BytesIO(r.content))
        assert doc.paragraphs[0].text == "Meeting minutes: standup" and doc.tables[0].cell(1, 3).text == "SEC-401"
        assert client.get(f"/api/meetings/standup/exports/{card['card_id']}.md").text.startswith("# Meeting minutes")
        assert client.get(f"/api/meetings/standup/exports/{card['card_id']}.exe").status_code == 404
        assert client.get("/api/meetings/standup/minutes/latest").json()["card"]["card_id"] == card["card_id"]
        chat = client.get("/api/meetings/standup").json()["chat"]
        assert chat[-1]["kind"] == "export" and chat[-1]["attachment"]["card_id"] == card["card_id"]


def test_writer_route():
    from app.agents import route

    for q in ("make the meeting minutes", "can you create a doc summarizing the spec", "prepare a report on last week's metrics",
              "write up the MoM", "export the notes as a word file"):
        assert route(q) == "writer", q
    assert route("draft an email to David with the minutes") == "email"
    assert route("summarize the meeting") is None
