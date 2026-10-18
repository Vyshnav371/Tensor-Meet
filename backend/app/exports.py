"""Downloadable documents (meeting minutes, reports) written by Tensor: Markdown -> Word (.docx) or Markdown (.md)."""
from __future__ import annotations

import io
import re

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def filename(title: str, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9 _.-]+", "", title).strip().replace(" ", "_")[:80] or "Tensor_document"
    return f"{base}.{ext}"


def _runs(paragraph, text: str) -> None:
    """Bold, italic, code and links inside a line."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("[") and "](" in part:
            label, url = part[1:-1].split("](", 1)
            paragraph.add_run(label)
            if url.startswith("http"):
                paragraph.add_run(f" ({url})").italic = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def to_docx(title: str, markdown: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = "Tensor"
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    lines = markdown.replace("\r\n", "\n").split("\n")
    i, paragraph = 0, []

    def flush():
        """Consecutive lines stay on their own lines (e.g. "**Date:**" then "**Attendees:**"), in one paragraph."""
        if paragraph:
            para = doc.add_paragraph()
            for n, text in enumerate(paragraph):
                if n:
                    para.add_run().add_break()
                _runs(para, text)
            paragraph.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            flush()
        elif m := re.match(r"^(#{1,4})\s+(.*)", stripped):
            flush()
            level = len(m.group(1))
            heading = doc.add_heading(level=0 if level == 1 and i == next((k for k, l in enumerate(lines) if l.strip()), 0) else min(level, 4))
            _runs(heading, m.group(2).strip())
        elif stripped.startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.fullmatch(r"\|?[\s:|-]+\|?", lines[i].strip()):  # skip |---|---| separators
                    block.append(_cells(lines[i]))
                i += 1
            width = max(len(r) for r in block)
            table = doc.add_table(rows=len(block), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(block):
                for c in range(width):
                    cell = table.cell(r, c)
                    cell.text = ""
                    _runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            continue
        elif m := re.match(r"^(\s*)[-*+]\s+(.*)", line):
            flush()
            indent = len(m.group(1).replace("\t", "  ")) // 2
            _runs(doc.add_paragraph(style="List Bullet 2" if indent else "List Bullet"), m.group(2))
        elif m := re.match(r"^\s*\d+[.)]\s+(.*)", line):
            flush()
            _runs(doc.add_paragraph(style="List Number"), m.group(1))
        elif stripped.startswith(">"):
            flush()
            quote = doc.add_paragraph()
            _runs(quote, stripped.lstrip("> "))
            for run in quote.runs:
                run.italic = True
        elif re.fullmatch(r"[-*_]{3,}", stripped):
            flush()
        else:
            paragraph.append(stripped)
        i += 1
    flush()
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
